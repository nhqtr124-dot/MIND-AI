"""Acceptance 6, 7, 8, 17, 18, 19 plus research and image editing."""

import io
import uuid
from typing import Any

import httpx
import pytest
import trimesh

from .conftest import FakeProvider, Session, jobs, setup_models, tiny_png


def _job(s: Session, job_id: str) -> dict[str, Any]:
    jobs()
    return s.get(f"/api/v1/jobs/{job_id}").json()


def _download(s: Session, artifact: dict[str, Any], fmt: str) -> tuple[str, bytes]:
    f = next(f for f in artifact["versions"][-1]["files"] if f["format"] == fmt)
    r = s.get(f"/api/v1/artifacts/{artifact['id']}/files/{f['name']}")
    assert r.status_code == 200
    return f["name"], r.content


def test_generate_and_download_docx(alice: Session) -> None:
    spec = {"title": "Team Charter", "blocks": [{"type": "heading", "text": "Mission"}, {"type": "paragraph", "text": "Build safe and reliable competition robots."}, {"type": "table", "columns": ["Role", "Name"], "rows": [["Lead", "Alice"]]}]}
    r = alice.post("/api/v1/documents/generate", json={"org_id": alice.org_id, "kind": "document", "format": "docx", "spec": spec})
    assert r.status_code == 202
    job = _job(alice, r.json()["job_id"])
    assert job["status"] == "completed", job
    art = alice.get(f"/api/v1/artifacts/{r.json()['artifact_id']}").json()
    assert art["status"] == "completed" and art["validation_status"] == "passed"
    name, data = _download(alice, art, "docx")
    assert name.endswith(".docx")
    from docx import Document

    doc = Document(io.BytesIO(data))
    assert "Team Charter" in [p.text for p in doc.paragraphs] and doc.tables[0].cell(1, 1).text == "Alice"
    # signed URL works without auth, and is bound to this file
    signed = alice.post(f"/api/v1/artifacts/{art['id']}/files/{name}/signed-url").json()
    token = signed["url"].split("/downloads/")[1]
    assert alice.client.get(f"/api/v1/downloads/{token}").content == data
    assert alice.client.get(f"/api/v1/downloads/{token[:-3]}abc").status_code == 403


def test_xlsx_formulas_validated(alice: Session) -> None:
    spec = {"sheets": [{"name": "BOM", "rows": [["Part", "Qty", "Unit", "Total"], ["Servo", 6, 12.5, "=B2*C2"], ["Frame", 1, 40, "=B3*C3"], ["Sum", None, None, "=SUM(D2:D3)"]]}], "expected": {"BOM!D4": 115}}
    r = alice.post("/api/v1/documents/generate", json={"org_id": alice.org_id, "kind": "spreadsheet", "format": "xlsx", "spec": spec})
    job = _job(alice, r.json()["job_id"])
    assert job["status"] == "completed"
    art = alice.get(f"/api/v1/artifacts/{r.json()['artifact_id']}").json()
    v = art["versions"][-1]["validation"]
    assert {c["id"]: c["status"] for c in v["checks"]}["expected_values"] == "pass" and v["metrics"]["computed"]["BOM!D4"] == 115
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(_download(alice, art, "xlsx")[1]))
    assert wb["BOM"]["D4"].value == "=SUM(D2:D3)"
    # a wrong expectation fails the job instead of returning a bad file as success
    spec["expected"] = {"BOM!D4": 999}
    r = alice.post("/api/v1/documents/generate", json={"org_id": alice.org_id, "kind": "spreadsheet", "format": "xlsx", "spec": spec})
    job = _job(alice, r.json()["job_id"])
    assert job["status"] == "failed" and "validation" in job["error"]["message"]
    assert alice.get(f"/api/v1/artifacts/{r.json()['artifact_id']}").json()["status"] == "failed"


def test_invalid_doc_spec_rejected(alice: Session) -> None:
    r = alice.post("/api/v1/documents/generate", json={"org_id": alice.org_id, "kind": "document", "format": "pptx", "spec": {"title": "x"}})
    assert r.status_code == 422


def test_arduino_holder_stl(alice: Session) -> None:
    proj = alice.post("/api/v1/projects", json={"org_id": alice.org_id, "name": "Robot"}).json()
    r = alice.post("/api/v1/cad/generate", json={"project_id": proj["id"], "template": "arduino_uno_holder", "params": {"rim_height": 8}, "formats": ["stl", "step", "3mf"]})
    assert r.status_code == 202
    job = _job(alice, r.json()["job_id"])
    assert job["status"] == "completed", job
    art = alice.get(f"/api/v1/artifacts/{r.json()['artifact_id']}").json()
    assert art["validation_status"] == "print_ready_checks_passed" and art["project_id"] == proj["id"]
    name, data = _download(alice, art, "stl")
    mesh = trimesh.load(io.BytesIO(data), file_type="stl")
    assert mesh.is_watertight and list(map(lambda v: round(v, 1), mesh.extents)) == [78.6, 63.3, 11.0]
    report = art["versions"][-1]["validation"]
    assert report["bom"][1]["item"].startswith("M3") and "disclaimer" in report["parts"][0]["validation"]
    assert {f["format"] for f in art["versions"][-1]["files"]} == {"stl", "step", "3mf", "json"}
    # new version with different parameters
    r = alice.post("/api/v1/cad/generate", json={"project_id": proj["id"], "template": "arduino_uno_holder", "params": {"margin": 8}, "artifact_id": art["id"]})
    _job(alice, r.json()["job_id"])
    art = alice.get(f"/api/v1/artifacts/{art['id']}").json()
    assert art["current_version"] == 2 and art["versions"][1]["params"]["margin"] == 8


def test_failed_cad_validation_is_not_print_ready(alice: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    import mind_cad.pipeline as pipeline
    from mind_cad.validation import Check, ValidationReport

    def failing(path: Any, **kw: Any) -> ValidationReport:
        return ValidationReport(file="x.stl", status="validation_failed", checks=[Check("watertight", "Watertight", "fail", "open edges detected")])

    monkeypatch.setattr(pipeline, "validate_mesh", failing)
    r = alice.post("/api/v1/cad/generate", json={"org_id": alice.org_id, "template": "standoff"})
    job = _job(alice, r.json()["job_id"])
    assert job["status"] == "completed" and job["result"]["status"] == "validation_failed"
    art = alice.get(f"/api/v1/artifacts/{r.json()['artifact_id']}").json()
    assert art["validation_status"] == "validation_failed" and art["status"] == "partially_completed"


def test_uploaded_broken_mesh_fails_validation(alice: Session) -> None:
    box = trimesh.creation.box(extents=(10, 10, 10))
    broken = trimesh.Trimesh(vertices=box.vertices, faces=box.faces[:-1], process=False)
    up = alice.post("/api/v1/files", data={"org_id": alice.org_id}, files={"file": ("broken.stl", broken.export(file_type="stl"), "model/stl")}).json()
    rep = alice.post("/api/v1/cad/validate", json={"file_id": up["id"]}).json()
    assert rep["status"] == "validation_failed" and not rep["passed"]


def test_cad_param_validation(alice: Session) -> None:
    r = alice.post("/api/v1/cad/generate", json={"org_id": alice.org_id, "template": "standoff", "params": {"outer_diameter": 3.5, "hole_diameter": 3.4}})
    assert r.status_code == 422
    assert alice.post("/api/v1/cad/generate", json={"org_id": alice.org_id, "template": "warp_drive"}).status_code == 404


def test_image_generation_produces_decodable_file(alice: Session, fake_provider: FakeProvider) -> None:
    """Mocked provider: proves the pipeline stores only decodable images; not proof of a live provider."""
    ids = setup_models(alice, fake_provider, ("gpt-image-fake",), gpt_image_fake={"price_per_image": "0.04"})
    r = alice.post("/api/v1/images/generate", json={"org_id": alice.org_id, "model_config_id": ids["gpt-image-fake"], "prompt": "a red robot", "n": 2})
    assert r.status_code == 202, r.text
    job = _job(alice, r.json()["job_id"])
    assert job["status"] == "completed"
    art = alice.get(f"/api/v1/artifacts/{r.json()['artifact_id']}").json()
    assert len(art["versions"][-1]["files"]) == 2
    from PIL import Image

    im = Image.open(io.BytesIO(_download(alice, art, "png")[1]))
    assert im.size == (8, 6)
    assert float(alice.get(f"/api/v1/admin/usage?org_id={alice.org_id}").json()["total_cost_usd"]) == pytest.approx(0.08)


def test_image_provider_garbage_is_rejected(alice: Session, fake_provider: FakeProvider, monkeypatch: pytest.MonkeyPatch) -> None:
    ids = setup_models(alice, fake_provider, ("gpt-image-fake",))
    orig = fake_provider.handler

    def bad(req: httpx.Request) -> httpx.Response:
        if req.url.path.endswith("/images/generations"):
            return httpx.Response(200, json={"data": [{"b64_json": "bm90IGFuIGltYWdl"}]})
        return orig(req)

    monkeypatch.setattr(fake_provider, "handler", bad)
    from mind_api.services import providers as ps

    ps.HTTP_CLIENT_FACTORY = lambda: httpx.AsyncClient(transport=httpx.MockTransport(bad))
    r = alice.post("/api/v1/images/generate", json={"org_id": alice.org_id, "model_config_id": ids["gpt-image-fake"], "prompt": "x"})
    job = _job(alice, r.json()["job_id"])
    assert job["status"] == "failed" and "not a valid image" in job["error"]["message"]
    art = alice.get(f"/api/v1/artifacts/{r.json()['artifact_id']}").json()
    assert art["status"] == "failed" and art["versions"] == []


def test_local_image_edit(alice: Session) -> None:
    up = alice.post("/api/v1/files", data={"org_id": alice.org_id}, files={"file": ("p.png", tiny_png(), "image/png")}).json()
    r = alice.post("/api/v1/images/edit", json={"source_file_id": up["id"], "ops": [{"op": "resize", "width": 16}, {"op": "rotate", "degrees": 90}, {"op": "grayscale"}], "output_format": "jpeg"})
    assert r.status_code == 200, r.text
    f = r.json()["versions"][0]["files"][0]
    assert (f["width"], f["height"], f["format"]) == (12, 16, "jpg")
    bad = alice.post("/api/v1/images/edit", json={"source_file_id": up["id"], "ops": [{"op": "crop", "width": 100, "height": 100}]})
    assert bad.status_code == 422


def test_unimplemented_studios_are_reported_honestly(alice: Session) -> None:
    caps = alice.get(f"/api/v1/capabilities?org_id={alice.org_id}").json()
    assert caps["features"]["video"]["status"] == "not_implemented"
    assert caps["features"]["voice"]["status"] == "not_implemented"
    assert caps["features"]["chat"]["status"] == "needs_configuration"
    assert caps["features"]["cad"]["status"] == "available"


def test_research_from_urls_with_cited_synthesis(alice: Session, fake_provider: FakeProvider, monkeypatch: pytest.MonkeyPatch) -> None:
    """Web pages and the model are mocked; checks staging, provenance and citation verification."""
    page = "<html><head><title>Servo basics</title></head><body><p>Hobby servos are controlled with a 50 Hz PWM signal where the pulse width sets the angle.</p><p>Most hobby servos rotate about 180 degrees in total.</p><p>Continuous rotation servos replace position control with speed control, so the same pulse width now sets direction and speed rather than angle.</p></body></html>"

    def web(req: httpx.Request) -> httpx.Response:
        if req.url.path == "/robots.txt":
            return httpx.Response(404)
        return httpx.Response(200, text=page, headers={"content-type": "text/html"})

    import mind_api.services.research_service as rs
    from mind_research import Fetcher

    monkeypatch.setattr(rs, "Fetcher", lambda: Fetcher(httpx.AsyncClient(transport=httpx.MockTransport(web)), check_ssrf=False))
    setup_models(alice, fake_provider)
    fake_provider.reply = 'Servos use "a 50 Hz PWM signal" for control [1]. They rotate roughly 180 degrees [1].'
    r = alice.post("/api/v1/research", json={"org_id": alice.org_id, "question": "How are hobby servos controlled?", "urls": ["https://servo.example/guide"]})
    job = _job(alice, r.json()["job_id"])
    assert job["status"] == "completed", job["error"]
    art = alice.get(f"/api/v1/artifacts/{r.json()['artifact_id']}").json()
    assert art["validation_status"] == "citations_verified"
    md = _download(alice, art, "md")[1].decode()
    assert "## Evidence (verbatim extracts)" in md and "https://servo.example/guide" in md and "retrieved" in md
    # a fabricated quote is caught
    fake_provider.reply = 'Servos use "a 400 Hz analog voltage" [1] according to [5].'
    r = alice.post("/api/v1/research", json={"org_id": alice.org_id, "question": "How are hobby servos controlled?", "urls": ["https://servo.example/guide"]})
    job = _job(alice, r.json()["job_id"])
    art = alice.get(f"/api/v1/artifacts/{r.json()['artifact_id']}").json()
    assert job["status"] == "partially_completed" and art["validation_status"] == "citations_failed"
    assert "Citation check failed" in _download(alice, art, "md")[1].decode()


def test_research_without_engine_or_urls_fails_clearly(alice: Session) -> None:
    r = alice.post("/api/v1/research", json={"org_id": alice.org_id, "question": "anything at all"})
    job = _job(alice, r.json()["job_id"])
    assert job["status"] == "failed" and "No web search provider" in job["error"]["message"]


def test_job_cancel(alice: Session) -> None:
    r = alice.post("/api/v1/cad/generate", json={"org_id": alice.org_id, "template": "standoff"})
    c = alice.post(f"/api/v1/jobs/{r.json()['job_id']}/cancel").json()
    assert c["status"] == "cancelled"
    jobs()
    assert alice.get(f"/api/v1/jobs/{r.json()['job_id']}").json()["status"] == "cancelled"
    assert alice.get(f"/api/v1/jobs/{uuid.uuid4()}").status_code == 404
