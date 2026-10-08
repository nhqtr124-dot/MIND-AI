"""Reopen generated files and check them against the spec that produced them."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Literal

from .spec import BulletList, DocumentSpec, Heading, Paragraph, PresentationSpec, Table, WorkbookSpec

Status = Literal["pass", "fail", "warn"]
_EXCEL_ERRORS = ("#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


@dataclass
class DocCheck:
    id: str
    status: Status
    detail: str


@dataclass
class DocValidation:
    file: str
    format: str
    checks: list[DocCheck] = field(default_factory=list)
    metrics: dict[str, Any] = field(default_factory=dict)

    @property
    def passed(self) -> bool:
        return all(c.status != "fail" for c in self.checks)

    def add(self, id: str, ok: bool, detail: str, warn_only: bool = False) -> None:
        self.checks.append(DocCheck(id, "pass" if ok else ("warn" if warn_only else "fail"), detail))

    def to_dict(self) -> dict[str, Any]:
        d = asdict(self)
        d["passed"] = self.passed
        return d


def _norm(s: str) -> str:
    # NFKC folds Arabic presentation forms (as extracted from shaped PDFs) back to base letters.
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", s)).strip().casefold()


def _expected_texts(spec: DocumentSpec) -> list[str]:
    out = [spec.title]
    for b in spec.blocks:
        if isinstance(b, Heading):
            out.append(b.text)
        elif isinstance(b, BulletList):
            out.extend(b.items)
        elif isinstance(b, Table):
            out.extend(b.columns)
    return out


def _word_coverage(spec: DocumentSpec, text: str) -> float:
    words = [
        w for b in spec.blocks if isinstance(b, Paragraph) for w in re.findall(r"\w+", b.text.casefold())
    ]
    if not words:
        return 1.0
    present = set(re.findall(r"\w+", text.casefold()))
    return sum(w in present for w in words) / len(words)


def validate_document(path: Path, fmt: str, spec: DocumentSpec) -> DocValidation:
    v = DocValidation(path.name, fmt)
    v.add(
        "exists",
        path.is_file() and path.stat().st_size > 0,
        f"{path.stat().st_size if path.exists() else 0} bytes",
    )
    if not v.passed:
        return v
    try:
        if fmt == "docx":
            from docx import Document

            d = Document(str(path))
            text = "\n".join(p.text for p in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    text += "\n" + " ".join(c.text for c in row.cells)
            headings = [
                p.text for p in d.paragraphs if p.style is not None and p.style.name.startswith("Heading")
            ]
            v.metrics.update(paragraphs=len(d.paragraphs), tables=len(d.tables), headings=len(headings))
            want = [b.text for b in spec.blocks if isinstance(b, Heading)]
            v.add(
                "headings",
                all(w in headings for w in want),
                f"{len(want)} expected headings, found {len(headings)}",
            )
            want_tables = sum(isinstance(b, Table) for b in spec.blocks)
            v.add("tables", len(d.tables) == want_tables, f"{len(d.tables)} tables, expected {want_tables}")
            v.add("opens", True, "reopened with python-docx")
        elif fmt == "pdf":
            import pymupdf

            with pymupdf.open(str(path)) as pdf:
                v.metrics["pages"] = pdf.page_count
                text = "\n".join(page.get_text() for page in pdf)
            v.add("opens", v.metrics["pages"] > 0, f"{v.metrics['pages']} pages, reopened with PyMuPDF")
        elif fmt in ("md", "html", "txt"):
            text = path.read_text(encoding="utf-8")
            if fmt == "html":
                v.add("html_structure", "<html" in text and "</html>" in text, "html root element present")
            v.add("opens", True, "decoded as UTF-8")
        else:
            v.add("format", False, f"no validator for {fmt}")
            return v
    except Exception as exc:  # noqa: BLE001 - any reopen failure is a validation failure
        v.add("opens", False, f"failed to reopen: {exc}")
        return v

    hay = _norm(text)
    missing = [t for t in _expected_texts(spec) if _norm(t) not in hay]
    v.add(
        "content",
        not missing,
        "all titles, headings, list items and table headers present"
        if not missing
        else f"missing: {missing[:5]}",
    )
    cov = _word_coverage(spec, unicodedata.normalize("NFKC", text))
    v.add("paragraph_text", cov >= 0.95, f"{cov:.0%} of paragraph words found")
    return v


def validate_pptx(path: Path, spec: PresentationSpec) -> DocValidation:
    v = DocValidation(path.name, "pptx")
    try:
        from pptx import Presentation

        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        v.add("opens", False, f"failed to reopen: {exc}")
        return v
    slides = list(prs.slides)
    v.add("opens", True, "reopened with python-pptx")
    v.metrics["slides"] = len(slides)
    v.add(
        "slide_count",
        len(slides) == len(spec.slides) + 1,
        f"{len(slides)} slides (title + {len(spec.slides)})",
    )
    titles = [s.shapes.title.text if s.shapes.title is not None else "" for s in slides[1:]]
    v.add("titles", titles == [s.title for s in spec.slides], "slide titles match the outline")
    all_text = " ".join(sh.text_frame.text for s in slides for sh in s.shapes if sh.has_text_frame)
    missing = [b for s in spec.slides for b in s.bullets if _norm(b) not in _norm(all_text)]
    v.add("bullets", not missing, "all bullet text present" if not missing else f"missing: {missing[:3]}")
    return v


def validate_xlsx(path: Path, spec: WorkbookSpec) -> DocValidation:
    v = DocValidation(path.name, "xlsx")
    try:
        from openpyxl import load_workbook

        wb = load_workbook(str(path))
    except Exception as exc:  # noqa: BLE001
        v.add("opens", False, f"failed to reopen: {exc}")
        return v
    v.add("opens", True, "reopened with openpyxl")
    v.add("sheets", wb.sheetnames == [s.name for s in spec.sheets], f"sheets {wb.sheetnames}")

    formulas: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas.append(f"{ws.title}!{c.coordinate}")
    v.metrics["formula_cells"] = len(formulas)
    v.metrics["charts"] = sum(len(ws._charts) for ws in wb.worksheets)

    if formulas or spec.expected:
        from pycel import ExcelCompiler

        xl = ExcelCompiler(filename=str(path))
        computed: dict[str, Any] = {}
        errors: list[str] = []
        for ref in formulas:
            try:
                val = xl.evaluate(ref)
            except Exception as exc:  # noqa: BLE001
                errors.append(f"{ref}: {exc}")
                continue
            computed[ref] = val
            if isinstance(val, str) and val in _EXCEL_ERRORS:
                errors.append(f"{ref} = {val}")
        v.add(
            "formulas_evaluate",
            not errors,
            f"{len(formulas)} formulas evaluated with pycel" if not errors else "; ".join(errors[:5]),
        )
        mismatches = []
        for ref, want in spec.expected.items():
            got = computed.get(ref)
            if got is None:
                try:
                    got = xl.evaluate(ref)
                except Exception as exc:  # noqa: BLE001
                    mismatches.append(f"{ref}: {exc}")
                    continue
            ok = (
                abs(float(got) - float(want)) < 1e-9 * max(1.0, abs(float(want)))
                if isinstance(want, (int, float))
                and not isinstance(want, bool)
                and isinstance(got, (int, float))
                else got == want
            )
            if not ok:
                mismatches.append(f"{ref}: expected {want!r}, computed {got!r}")
        if spec.expected:
            v.add(
                "expected_values",
                not mismatches,
                f"{len(spec.expected)} expected values match"
                if not mismatches
                else "; ".join(mismatches[:5]),
            )
        v.metrics["computed"] = {k: computed[k] for k in list(computed)[:50]}
    return v


def validate_csv(path: Path, columns: list[str], n_rows: int) -> DocValidation:
    import csv

    v = DocValidation(path.name, "csv")
    with path.open(newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    v.add("header", bool(rows) and rows[0] == columns, "header row matches")
    v.add("rows", len(rows) - 1 == n_rows, f"{len(rows) - 1} data rows, expected {n_rows}")
    return v
