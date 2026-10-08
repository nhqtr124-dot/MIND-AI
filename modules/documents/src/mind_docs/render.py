"""Renderers: spec -> real file on disk."""

from __future__ import annotations

import csv
import html
import io
import os
from pathlib import Path
from typing import Any

from .spec import (
    BulletList,
    Chart,
    Code,
    DocumentSpec,
    Heading,
    PageBreak,
    Paragraph,
    PresentationSpec,
    Table,
    WorkbookSpec,
)

_RTL_LANGS = {"ar", "he", "fa", "ur"}
_FONT_CANDIDATES = [
    os.environ.get("MIND_PDF_FONT", ""),
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
]


def chart_png(chart: Chart, width_in: float = 6.0, height_in: float = 3.2) -> bytes:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(width_in, height_in), dpi=150)
    try:
        if chart.kind == "pie":
            ax.pie(chart.series[0].values, labels=chart.labels, autopct="%1.0f%%")
            ax.axis("equal")
        elif chart.kind == "line":
            for s in chart.series:
                ax.plot(chart.labels, s.values, marker="o", label=s.name)
        else:
            n = len(chart.series)
            w = 0.8 / n
            xs = range(len(chart.labels))
            for i, s in enumerate(chart.series):
                ax.bar([x + i * w - 0.4 + w / 2 for x in xs], s.values, width=w, label=s.name)
            ax.set_xticks(list(xs), chart.labels)
        if chart.title:
            ax.set_title(chart.title)
        if chart.kind != "pie" and len(chart.series) > 1:
            ax.legend()
        ax.spines[["top", "right"]].set_visible(False)
        fig.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format="png")
        return buf.getvalue()
    finally:
        plt.close(fig)


def _cell(v: Any) -> str:
    return "" if v is None else str(v)


# ---------------------------------------------------------------------------- DOCX


def render_docx(spec: DocumentSpec, path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    rtl = spec.language in _RTL_LANGS
    doc = Document()
    doc.core_properties.title = spec.title
    if spec.author:
        doc.core_properties.author = spec.author

    def mark_rtl(par: Any) -> None:
        if rtl:
            ppr = par._p.get_or_add_pPr()
            ppr.append(ppr.makeelement(qn("w:bidi"), {}))
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    mark_rtl(doc.add_heading(spec.title, level=0))
    if spec.subtitle:
        mark_rtl(doc.add_paragraph(spec.subtitle, style="Subtitle"))
    for b in spec.blocks:
        if isinstance(b, Heading):
            mark_rtl(doc.add_heading(b.text, level=b.level))
        elif isinstance(b, Paragraph):
            mark_rtl(doc.add_paragraph(b.text))
        elif isinstance(b, BulletList):
            for it in b.items:
                mark_rtl(doc.add_paragraph(it, style="List Number" if b.ordered else "List Bullet"))
        elif isinstance(b, Table):
            t = doc.add_table(rows=1 + len(b.rows), cols=len(b.columns))
            t.style = "Light Grid Accent 1"
            for j, c in enumerate(b.columns):
                t.cell(0, j).text = c
            for i, row in enumerate(b.rows, start=1):
                for j, v in enumerate(row):
                    t.cell(i, j).text = _cell(v)
            if b.caption:
                cap = doc.add_paragraph(b.caption, style="Caption")
                mark_rtl(cap)
        elif isinstance(b, Chart):
            doc.add_picture(io.BytesIO(chart_png(b)), width=Inches(6))
            if b.title:
                doc.add_paragraph(b.title, style="Caption")
        elif isinstance(b, Code):
            p = doc.add_paragraph()
            run = p.add_run(b.text)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif isinstance(b, PageBreak):
            doc.add_page_break()
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------- PDF


def _pdf_font() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for cand in _FONT_CANDIDATES:
        if cand and Path(cand).is_file():
            bold = cand.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
            if "MindSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("MindSans", cand))
                pdfmetrics.registerFont(TTFont("MindSans-Bold", bold if Path(bold).is_file() else cand))
            return "MindSans", "MindSans-Bold"
    return "Helvetica", "Helvetica-Bold"


def _shape(text: str, rtl: bool) -> str:
    if not rtl:
        return html.escape(text)
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display

        return html.escape(get_display(arabic_reshaper.reshape(text)))
    except ImportError:
        return html.escape(text)


def render_pdf(spec: DocumentSpec, path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Image,
        ListFlowable,
        ListItem,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
        TableStyle,
    )
    from reportlab.platypus import PageBreak as RLPageBreak
    from reportlab.platypus import Paragraph as RLParagraph
    from reportlab.platypus import Table as RLTable

    font, bold = _pdf_font()
    rtl = spec.language in _RTL_LANGS
    ss = getSampleStyleSheet()
    align = {"alignment": TA_RIGHT} if rtl else {}
    body = ParagraphStyle("body", parent=ss["BodyText"], fontName=font, fontSize=10.5, leading=15, **align)
    title = ParagraphStyle("title", parent=ss["Title"], fontName=bold, **align)
    sub = ParagraphStyle("sub", parent=ss["Heading3"], fontName=font, textColor=colors.grey, **align)
    hs = {
        lvl: ParagraphStyle(f"h{lvl}", parent=ss[f"Heading{min(lvl, 4)}"], fontName=bold, **align)
        for lvl in range(1, 5)
    }

    story: list[Any] = [RLParagraph(_shape(spec.title, rtl), title)]
    if spec.subtitle:
        story.append(RLParagraph(_shape(spec.subtitle, rtl), sub))
    story.append(Spacer(1, 0.4 * cm))
    for b in spec.blocks:
        if isinstance(b, Heading):
            story.append(RLParagraph(_shape(b.text, rtl), hs[b.level]))
        elif isinstance(b, Paragraph):
            story.append(RLParagraph(_shape(b.text, rtl), body))
        elif isinstance(b, BulletList):
            story.append(
                ListFlowable(
                    [ListItem(RLParagraph(_shape(i, rtl), body)) for i in b.items],
                    bulletType="1" if b.ordered else "bullet",
                    bulletFontName=font,
                )
            )
        elif isinstance(b, Table):
            data = [[RLParagraph(_shape(c, rtl), body) for c in b.columns]] + [
                [RLParagraph(_shape(_cell(v), rtl), body) for v in r] for r in b.rows
            ]
            t = RLTable(data, repeatRows=1, hAlign="RIGHT" if rtl else "LEFT")
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF0FF")),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8BCD9")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(t)
            if b.caption:
                story.append(RLParagraph(_shape(b.caption, rtl), sub))
        elif isinstance(b, Chart):
            story.append(Image(io.BytesIO(chart_png(b)), width=16 * cm, height=8.5 * cm))
        elif isinstance(b, Code):
            story.append(
                Preformatted(b.text, ParagraphStyle("code", fontName="Courier", fontSize=8.5, leading=11))
            )
        elif isinstance(b, PageBreak):
            story.append(RLPageBreak())
        story.append(Spacer(1, 0.2 * cm))

    doc = SimpleDocTemplate(str(path), pagesize=A4, title=spec.title, author=spec.author or "MIND AI")
    doc.build(story)
    return path


# ---------------------------------------------------------------------------- Markdown / HTML / TXT


def to_markdown(spec: DocumentSpec) -> str:
    out = [f"# {spec.title}", ""]
    if spec.subtitle:
        out += [f"_{spec.subtitle}_", ""]
    for b in spec.blocks:
        if isinstance(b, Heading):
            out += [f"{'#' * (b.level + 1)} {b.text}", ""]
        elif isinstance(b, Paragraph):
            out += [b.text, ""]
        elif isinstance(b, BulletList):
            out += [f"{i + 1}. {t}" if b.ordered else f"- {t}" for i, t in enumerate(b.items)] + [""]
        elif isinstance(b, Table):
            out.append("| " + " | ".join(b.columns) + " |")
            out.append("|" + "---|" * len(b.columns))
            out += ["| " + " | ".join(_cell(v).replace("|", "\\|") for v in r) + " |" for r in b.rows]
            out += ([f"_{b.caption}_"] if b.caption else []) + [""]
        elif isinstance(b, Chart):
            out.append(f"**Chart: {b.title or b.kind}**")
            out.append("| Label | " + " | ".join(s.name for s in b.series) + " |")
            out.append("|---|" + "---|" * len(b.series))
            for i, label in enumerate(b.labels):
                out.append(f"| {label} | " + " | ".join(str(s.values[i]) for s in b.series) + " |")
            out.append("")
        elif isinstance(b, Code):
            out += [f"```{b.language}", b.text, "```", ""]
        elif isinstance(b, PageBreak):
            out += ["---", ""]
    return "\n".join(out).rstrip() + "\n"


def to_html(spec: DocumentSpec) -> str:
    from markdown_it import MarkdownIt

    md = MarkdownIt("commonmark").enable("table")
    direction = "rtl" if spec.language in _RTL_LANGS else "ltr"
    body = md.render(to_markdown(spec))
    return (
        f'<!doctype html>\n<html lang="{html.escape(spec.language)}" dir="{direction}"><head><meta charset="utf-8">'
        f"<title>{html.escape(spec.title)}</title>"
        "<style>body{font-family:system-ui,sans-serif;max-width:52rem;margin:2rem auto;padding:0 1rem;line-height:1.6}"
        "table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:.3rem .6rem}</style></head>"
        f"<body>\n{body}</body></html>\n"
    )


def to_text(spec: DocumentSpec) -> str:
    out = [spec.title, "=" * min(len(spec.title), 80), ""]
    if spec.subtitle:
        out += [spec.subtitle, ""]
    for b in spec.blocks:
        if isinstance(b, Heading):
            out += [b.text, "-" * min(len(b.text), 80)]
        elif isinstance(b, Paragraph):
            out += [b.text, ""]
        elif isinstance(b, BulletList):
            out += [f"{i + 1}. {t}" if b.ordered else f"* {t}" for i, t in enumerate(b.items)] + [""]
        elif isinstance(b, Table):
            out += ["\t".join(b.columns)] + ["\t".join(_cell(v) for v in r) for r in b.rows] + [""]
        elif isinstance(b, Chart):
            out += (
                [f"[{b.kind} chart] {b.title}"]
                + [
                    f"{label}: " + ", ".join(f"{s.name}={s.values[i]}" for s in b.series)
                    for i, label in enumerate(b.labels)
                ]
                + [""]
            )
        elif isinstance(b, Code):
            out += [b.text, ""]
    return "\n".join(out).rstrip() + "\n"


def render_document(spec: DocumentSpec, fmt: str, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "docx":
        return render_docx(spec, path)
    if fmt == "pdf":
        return render_pdf(spec, path)
    if fmt == "md":
        path.write_text(to_markdown(spec), encoding="utf-8")
    elif fmt == "html":
        path.write_text(to_html(spec), encoding="utf-8")
    elif fmt == "txt":
        path.write_text(to_text(spec), encoding="utf-8")
    else:
        raise ValueError(f"unsupported document format '{fmt}'")
    return path


# ---------------------------------------------------------------------------- PPTX


def render_pptx(spec: PresentationSpec, path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    s0.shapes.title.text = spec.title
    if spec.subtitle and len(s0.placeholders) > 1:
        s0.placeholders[1].text = spec.subtitle
    for sl in spec.slides:
        layout = prs.slide_layouts[1] if sl.bullets and not (sl.table or sl.chart) else prs.slide_layouts[5]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = sl.title
        top = Inches(1.6)
        if sl.bullets:
            if layout == prs.slide_layouts[1]:
                tf = s.placeholders[1].text_frame
            else:
                tb = s.shapes.add_textbox(Inches(0.7), top, Inches(5.5), Inches(5))
                tf = tb.text_frame
                tf.word_wrap = True
            tf.text = sl.bullets[0]
            for b in sl.bullets[1:]:
                tf.add_paragraph().text = b
        left = Inches(6.6) if sl.bullets else Inches(0.7)
        width = Inches(6.1) if sl.bullets else Inches(11.9)
        if sl.table:
            t = sl.table
            shape = s.shapes.add_table(
                1 + len(t.rows), len(t.columns), left, top, width, Inches(0.4) * (1 + len(t.rows))
            )
            for j, c in enumerate(t.columns):
                shape.table.cell(0, j).text = c
            for i, r in enumerate(t.rows, start=1):
                for j, v in enumerate(r):
                    cell = shape.table.cell(i, j)
                    cell.text = _cell(v)
                    cell.text_frame.paragraphs[0].runs[0].font.size = Pt(12) if cell.text else None  # type: ignore[assignment]
        elif sl.chart:
            s.shapes.add_picture(io.BytesIO(chart_png(sl.chart, 7, 4)), left, top, width=width)
        if sl.notes:
            s.notes_slide.notes_text_frame.text = sl.notes
    prs.core_properties.title = spec.title
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


# ---------------------------------------------------------------------------- XLSX / CSV


def render_xlsx(spec: WorkbookSpec, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font
    from openpyxl.utils.cell import range_boundaries

    wb = Workbook()
    wb.remove(wb.active)
    for sh in spec.sheets:
        ws = wb.create_sheet(sh.name)
        for row in sh.rows:
            ws.append(list(row))
        if sh.bold_header and sh.rows:
            for c in ws[1]:
                c.font = Font(bold=True)
        for col, w in sh.column_widths.items():
            ws.column_dimensions[col].width = w
        for ch in sh.charts:
            chart = BarChart() if ch.kind == "bar" else LineChart()
            chart.title = ch.title or None
            c1, r1, c2, r2 = range_boundaries(ch.data_range)
            data = Reference(ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2)
            k1, kr1, k2, kr2 = range_boundaries(ch.categories_range)
            cats = Reference(ws, min_col=k1, min_row=kr1, max_col=k2, max_row=kr2)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, ch.anchor)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def render_csv(columns: list[str], rows: list[list[Any]], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(columns)
        w.writerows([[_cell(v) for v in r] for r in rows])
    return path
