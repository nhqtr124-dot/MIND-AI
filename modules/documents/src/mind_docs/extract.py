"""Text extraction from uploaded files (for Q&A and memory indexing)."""

from __future__ import annotations

import csv
import io
import shutil
import subprocess
import tempfile
from html.parser import HTMLParser
from pathlib import Path

TEXT_TYPES = {".txt", ".md", ".markdown", ".json", ".py", ".ts", ".tsx", ".js", ".css", ".yaml", ".yml", ".toml", ".sql"}
SUPPORTED = TEXT_TYPES | {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".html", ".htm"}


class ExtractionError(ValueError):
    pass


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style", "noscript"):
            self._skip += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style", "noscript") and self._skip:
            self._skip -= 1
        if tag in ("p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr"):
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def html_to_text(markup: str) -> str:
    p = _HTMLText()
    p.feed(markup)
    lines = [ln.strip() for ln in "".join(p.parts).splitlines()]
    return "\n".join(ln for ln in lines if ln)


def extract_text(path: Path, max_chars: int = 2_000_000) -> str:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ExtractionError(f"text extraction is not supported for '{ext}' files")
    if ext == ".pdf":
        import pymupdf

        with pymupdf.open(str(path)) as pdf:
            text = "\n\n".join(f"[page {i + 1}]\n{page.get_text()}" for i, page in enumerate(pdf))
        if not text.replace("[page", "").strip(" \n0123456789]"):
            raise ExtractionError("PDF has no extractable text layer (scanned?); OCR is not installed")
    elif ext == ".docx":
        from docx import Document

        d = Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
        text = "\n".join(parts)
    elif ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(path))
        text = "\n\n".join(
            f"[slide {i + 1}]\n" + "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
            for i, s in enumerate(prs.slides)
        )
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(str(path), data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"[sheet {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    out.append("\t".join("" if v is None else str(v) for v in row))
        text = "\n".join(out)
    elif ext == ".csv":
        with path.open(newline="", encoding="utf-8", errors="replace") as f:
            text = "\n".join("\t".join(r) for r in csv.reader(f))
    elif ext in (".html", ".htm"):
        text = html_to_text(path.read_text(encoding="utf-8", errors="replace"))
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    return text[:max_chars]


def libreoffice_available() -> bool:
    return shutil.which("soffice") is not None


def convert_with_libreoffice(src: Path, target_ext: str, out_dir: Path, timeout: int = 120) -> Path:
    """Convert office documents with LibreOffice headless (e.g. docx -> pdf)."""
    soffice = shutil.which("soffice")
    if not soffice:
        raise ExtractionError("LibreOffice is not installed; document conversion is unavailable")
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="mind-lo-") as profile:
        proc = subprocess.run(
            [
                soffice,
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--convert-to",
                target_ext,
                "--outdir",
                str(out_dir),
                str(src),
            ],
            capture_output=True,
            timeout=timeout,
            check=False,
        )
    out = out_dir / f"{src.stem}.{target_ext}"
    if proc.returncode != 0 or not out.is_file():
        raise ExtractionError(f"conversion failed: {proc.stderr.decode(errors='replace')[-500:]}")
    return out


def render_pdf_previews(pdf: Path, max_pages: int = 3, zoom: float = 1.0) -> list[bytes]:
    import pymupdf

    out = []
    with pymupdf.open(str(pdf)) as doc:
        for page in list(doc)[:max_pages]:
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
            out.append(pix.tobytes("png"))
    return out


def csv_bytes(columns: list[str], rows: list[list[object]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    w.writerows(rows)
    return buf.getvalue().encode()
